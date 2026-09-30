import os
from typing import List, Dict, Any, Tuple, Optional
import pypdf
import docx

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:
    try:
        from langchain.text_splitter import RecursiveCharacterTextSplitter
    except ImportError:
        RecursiveCharacterTextSplitter = None


DEFAULT_CHUNK_SIZE = int(os.environ.get("CHUNK_SIZE", 800))
DEFAULT_CHUNK_OVERLAP = int(os.environ.get("CHUNK_OVERLAP", 120))
MAX_DOCUMENT_PAGES = int(os.environ.get("MAX_DOCUMENT_PAGES", 100))
MAX_DOCUMENT_CHUNKS = int(os.environ.get("MAX_DOCUMENT_CHUNKS", 500))

class DocumentProcessor:
    """
    Extracts text and page metadata from PDF, DOCX, and TXT files,
    and applies recursive character chunking with overlap.
    Includes memory/resource protection against oversized documents.
    """
    def __init__(
        self,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
        max_pages: Optional[int] = None,
        max_chunks: Optional[int] = None
    ):
        self.chunk_size = chunk_size or DEFAULT_CHUNK_SIZE
        self.chunk_overlap = chunk_overlap or DEFAULT_CHUNK_OVERLAP
        self.max_pages = max_pages or MAX_DOCUMENT_PAGES
        self.max_chunks = max_chunks or MAX_DOCUMENT_CHUNKS

        if RecursiveCharacterTextSplitter:
            self.splitter = RecursiveCharacterTextSplitter(
                chunk_size=self.chunk_size,
                chunk_overlap=self.chunk_overlap,
                separators=["\n\n", "\n", ". ", " ", ""]
            )
        else:
            self.splitter = None

    def _fallback_split(self, text: str) -> List[str]:
        """Pure Python fallback chunker if LangChain splitter is unavailable."""
        if not text:
            return []
        chunks = []
        start = 0
        text_len = len(text)
        while start < text_len:
            end = min(start + self.chunk_size, text_len)
            chunks.append(text[start:end])
            if end == text_len:
                break
            start += max(1, self.chunk_size - self.chunk_overlap)
        return chunks

    def split_text(self, text: str) -> List[str]:
        if self.splitter:
            return self.splitter.split_text(text)
        return self._fallback_split(text)

    def extract_from_pdf(self, file_path: str) -> Tuple[List[Dict[str, Any]], int]:
        """
        Extracts text from PDF page by page.
        Returns (list_of_pages, total_pages) where each page has {"page_number": int, "text": str}.
        """
        pages_data = []
        reader = pypdf.PdfReader(file_path)
        total_pages = len(reader.pages)

        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                pages_data.append({"page_number": i + 1, "text": text})

        return pages_data, total_pages

    def extract_from_docx(self, file_path: str) -> Tuple[List[Dict[str, Any]], int]:
        """
        Extracts paragraphs and tables from DOCX.
        Returns (list_of_sections, total_sections).
        """
        doc = docx.Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        if not full_text:
            return [], 1

        # Treat docx sections/pages
        return [{"page_number": 1, "text": full_text}], 1

    def extract_from_text(self, file_path: str) -> Tuple[List[Dict[str, Any]], int]:
        """Extracts text from TXT, MD, CSV, etc. with encoding fallbacks."""
        encodings = ['utf-8', 'latin-1', 'cp1252']
        text = ""
        for enc in encodings:
            try:
                with open(file_path, 'r', encoding=enc) as f:
                    text = f.read()
                break
            except (UnicodeDecodeError, LookupError):
                continue

        text = text.strip()
        if not text:
            return [], 1
        return [{"page_number": 1, "text": text}], 1

    def process_file(
        self,
        file_path: str,
        filename: str,
        document_id: int,
        user_id: str
    ) -> Tuple[List[Dict[str, Any]], int]:
        """
        Processes a file, extracts text, chunks it, and returns chunks + total pages.
        Each chunk:
        {
            "content": str,
            "page_number": int,
            "chunk_index": int,
            "document_id": int,
            "document_name": str,
            "user_id": str,
            "metadata": dict
        }
        """
        ext = os.path.splitext(filename)[1].lower()

        if ext == '.pdf':
            pages_data, total_pages = self.extract_from_pdf(file_path)
        elif ext in ('.docx', '.doc'):
            pages_data, total_pages = self.extract_from_docx(file_path)
        else:
            pages_data, total_pages = self.extract_from_text(file_path)

        if total_pages > self.max_pages:
            raise ValueError(
                f"Document exceeds maximum page limit of {self.max_pages} pages (found {total_pages})."
            )

        all_chunks = []
        chunk_idx = 0

        for page in pages_data:
            page_num = page["page_number"]
            page_text = page["text"]
            split_pieces = self.split_text(page_text)

            for piece in split_pieces:
                piece_clean = piece.strip()
                if len(piece_clean) < 15:  # Skip tiny whitespace/artifacts
                    continue

                if chunk_idx >= self.max_chunks:
                    raise ValueError(
                        f"Document exceeds maximum limit of {self.max_chunks} chunks."
                    )

                metadata = {
                    "document_id": document_id,
                    "document_name": filename,
                    "page_number": page_num,
                    "chunk_index": chunk_idx,
                    "user_id": user_id
                }

                all_chunks.append({
                    "content": piece_clean,
                    "page_number": page_num,
                    "chunk_index": chunk_idx,
                    "document_id": document_id,
                    "document_name": filename,
                    "user_id": user_id,
                    "metadata": metadata
                })
                chunk_idx += 1

        return all_chunks, max(total_pages, 1)
